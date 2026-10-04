import io
import logging
from datetime import datetime
from typing import Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.services.analytics_service import AnalyticsService

logger = logging.getLogger("sales_dashboard.export")


class ExportService:
    """Service to export filtered analytical reports in CSV and professional multi-tab Excel workbooks."""

    def __init__(self, analytics_service: AnalyticsService = None):
        self.analytics = analytics_service or AnalyticsService()

    def export_filtered_csv(self, filters: Dict[str, Any], max_rows: int = 50000) -> str:
        """Exports filtered order records to a CSV string."""
        where_clause, params = self.analytics._build_filter_clause(filters)
        params["limit"] = max_rows

        sql = f"""
        SELECT
            o.order_number AS "Order ID",
            o.order_date AS "Order Date",
            c.customer_code AS "Customer Code",
            c.first_name || ' ' || c.last_name AS "Customer Name",
            c.city AS "City",
            c.country AS "Country",
            r.region_name AS "Region",
            p.product_code AS "Product Code",
            p.product_name AS "Product Name",
            cat.category_name AS "Category",
            oi.quantity AS "Quantity",
            oi.unit_price AS "Unit Price",
            oi.discount_rate AS "Discount Rate",
            oi.total_price AS "Line Total",
            o.payment_method AS "Payment Method",
            o.order_status AS "Status"
        FROM orders o
        JOIN order_items oi ON o.order_id = oi.order_id
        JOIN products p ON oi.product_id = p.product_id
        JOIN categories cat ON p.category_id = cat.category_id
        JOIN customers c ON o.customer_id = c.customer_id
        JOIN regions r ON o.region_id = r.region_id
        WHERE {where_clause}
        ORDER BY o.order_date DESC, o.order_id DESC
        LIMIT :limit
        """

        with self.analytics.engine.connect() as conn:
            from sqlalchemy import text
            rows = conn.execute(text(sql), params).mappings().fetchall()
            df = pd.DataFrame([dict(r) for r in rows])

        output = io.StringIO()
        df.to_csv(output, index=False)
        return output.getvalue()

    def export_excel_report(self, filters: Dict[str, Any]) -> io.BytesIO:
        """
        Creates an executive multi-sheet Excel workbook with styled headers,
        formatted currency, KPI cards, and analytical breakdowns.
        """
        kpis = self.analytics.get_kpis(filters)
        monthly_trend = self.analytics.get_monthly_sales_trend(filters)
        top_prods = self.analytics.get_top_products(filters, limit=25)
        regional = self.analytics.get_region_performance(filters)
        segments = self.analytics.get_customer_segments(filters)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style definitions
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        sub_header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        card_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        accent_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
        card_lbl_font = Font(name="Calibri", size=9, bold=True, color="64748B")
        card_val_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        # ----------------------------------------------------
        # Sheet 1: Executive KPI Summary
        # ----------------------------------------------------
        ws1 = wb.create_sheet(title="Executive Summary")
        ws1.views.sheetView[0].showGridLines = True

        ws1["A1"] = "Sales Performance Analytics - Executive Summary"
        ws1["A1"].font = title_font
        ws1["A2"] = f"Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Active Filters: Date Range: {filters.get('start_date', 'All')} to {filters.get('end_date', 'All')} | Region: {filters.get('region_id', 'All')} | Category: {filters.get('category_id', 'All')}"
        ws1["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

        # KPI Tiles
        kpi_metrics = [
            ("TOTAL REVENUE", f"${kpis['total_revenue']:,.2f}", "Units Sold", f"{kpis['units_sold']:,}"),
            ("TOTAL ORDERS", f"{kpis['total_orders']:,}", "Avg Order Value (AOV)", f"${kpis['aov']:,.2f}"),
            ("TOTAL CUSTOMERS", f"{kpis['total_customers']:,}", "YoY Revenue Growth", f"{kpis['growth_pct']:+.2f}%"),
            ("TOP PERFORMING REGION", kpis['top_region']['name'], "Best-Selling Product", kpis['best_product']['name'])
        ]

        row_cursor = 4
        for row_data in kpi_metrics:
            # Card 1 (A-C)
            ws1.cell(row=row_cursor, column=1, value=row_data[0]).font = card_lbl_font
            ws1.cell(row=row_cursor + 1, column=1, value=row_data[1]).font = card_val_font
            # Card 2 (D-F)
            ws1.cell(row=row_cursor, column=4, value=row_data[2]).font = card_lbl_font
            ws1.cell(row=row_cursor + 1, column=4, value=row_data[3]).font = card_val_font
            row_cursor += 3

        # ----------------------------------------------------
        # Sheet 2: Monthly Sales Trend
        # ----------------------------------------------------
        ws2 = wb.create_sheet(title="Monthly Trends")
        ws2.views.sheetView[0].showGridLines = True
        ws2.append(["Monthly Sales Performance"])
        ws2["A1"].font = title_font
        ws2.append([])

        m_headers = ["Year-Month", "Total Revenue ($)", "Total Orders", "Units Sold", "AOV ($)"]
        ws2.append(m_headers)
        header_row = ws2[3]
        for cell in header_row:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for m in monthly_trend:
            ws2.append([
                m["year_month"],
                m["revenue"],
                m["orders_count"],
                m["units_sold"],
                m["aov"]
            ])

        for row in ws2.iter_rows(min_row=4, max_col=5):
            row[0].alignment = Alignment(horizontal="center")
            row[1].number_format = "$#,##0.00"
            row[2].number_format = "#,##0"
            row[3].number_format = "#,##0"
            row[4].number_format = "$#,##0.00"
            for cell in row:
                cell.border = thin_border
                cell.font = regular_font

        # ----------------------------------------------------
        # Sheet 3: Product Performance
        # ----------------------------------------------------
        ws3 = wb.create_sheet(title="Product Performance")
        ws3.views.sheetView[0].showGridLines = True
        ws3.append(["Top Product Performance Analysis"])
        ws3["A1"].font = title_font
        ws3.append([])

        p_headers = ["Rank", "Product Code", "Product Name", "Category", "Unit Price ($)", "Units Sold", "Total Revenue ($)"]
        ws3.append(p_headers)
        for cell in ws3[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, p in enumerate(top_prods, start=1):
            ws3.append([
                idx,
                p["product_code"],
                p["product_name"],
                p["category_name"],
                p["unit_price"],
                p["units_sold"],
                p["total_revenue"]
            ])

        for row in ws3.iter_rows(min_row=4, max_col=7):
            row[0].alignment = Alignment(horizontal="center")
            row[4].number_format = "$#,##0.00"
            row[5].number_format = "#,##0"
            row[6].number_format = "$#,##0.00"
            for cell in row:
                cell.border = thin_border
                cell.font = regular_font

        # ----------------------------------------------------
        # Sheet 4: Regional Analysis
        # ----------------------------------------------------
        ws4 = wb.create_sheet(title="Regional Analysis")
        ws4.views.sheetView[0].showGridLines = True
        ws4.append(["Territory Sales & Regional Rankings"])
        ws4["A1"].font = title_font
        ws4.append([])

        r_headers = ["Rank", "Region Name", "Region Code", "Country", "Total Orders", "Units Sold", "Total Revenue ($)", "Revenue Share %", "AOV ($)"]
        ws4.append(r_headers)
        for cell in ws4[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r in regional:
            ws4.append([
                r["rank"],
                r["region_name"],
                r["region_code"],
                r["country"],
                r["total_orders"],
                r["units_sold"],
                r["total_revenue"],
                r["share_pct"] / 100.0,
                r["aov"]
            ])

        for row in ws4.iter_rows(min_row=4, max_col=9):
            row[0].alignment = Alignment(horizontal="center")
            row[4].number_format = "#,##0"
            row[5].number_format = "#,##0"
            row[6].number_format = "$#,##0.00"
            row[7].number_format = "0.0%"
            row[8].number_format = "$#,##0.00"
            for cell in row:
                cell.border = thin_border
                cell.font = regular_font

        # ----------------------------------------------------
        # Sheet 5: Customer Segments
        # ----------------------------------------------------
        ws5 = wb.create_sheet(title="Customer Segments")
        ws5.views.sheetView[0].showGridLines = True
        ws5.append(["Customer Spending Segments Analysis"])
        ws5["A1"].font = title_font
        ws5.append([])

        cs_headers = ["Segment Name", "Customer Count", "Share %", "Total Orders", "Total Segment Spend ($)", "Avg Spend / Customer ($)"]
        ws5.append(cs_headers)
        for cell in ws5[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for s in segments.get("segments", []):
            ws5.append([
                s["segment_name"],
                s["customer_count"],
                s["share_pct"] / 100.0,
                s["total_orders"],
                s["total_segment_spend"],
                s["avg_spend_per_customer"]
            ])

        for row in ws5.iter_rows(min_row=4, max_col=6):
            row[1].number_format = "#,##0"
            row[2].number_format = "0.0%"
            row[3].number_format = "#,##0"
            row[4].number_format = "$#,##0.00"
            row[5].number_format = "$#,##0.00"
            for cell in row:
                cell.border = thin_border
                cell.font = regular_font

        # Auto-adjust column widths across all sheets
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    max_len = max(max_len, len(val_str))
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return file_stream
